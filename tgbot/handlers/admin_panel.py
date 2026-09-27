"""
Админка прямо в боте — /admin или кнопка "🛠 Admin panel".
Доступ: TGUser.is_admin=True или tg_id из ADMIN_IDS в .env.
Все тексты — на языке админа (tgbot/i18n.py, ключи adm_*).

Что умеет:
  📊 Статистика   — тот же отчёт, что приходит каждый вечер
  📅 Мероприятия  — список → карточка:
                      📥 Excel участников (для сертификатов)
                      ✅ список пришедших
                      ➕ добавить участника (поиск по @username/телефону/имени/email/ID)
                      ✉️ сообщение всем участникам мероприятия
  🔎 Поиск        — карточка человека: роль (с поздравлением), админка, добавить в мероприятие
  📥 Excel        — вся база
  📢 Рассылка     — подсказка по командам /send, /regionsend и т.д.

Плюс здесь же ловим блокировку бота юзером (my_chat_member) для статистики.
"""
import asyncio
import logging
import re
from html import escape
from io import BytesIO

from aiogram import types, Dispatcher
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import State, StatesGroup
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.utils import exceptions
from asgiref.sync import sync_to_async
from django.db.models import Count, Q
from django.utils import timezone

from app_telegram.models import TGUser, EcoProject, ProjectParticipation
from tgbot.i18n import (
    t, variants, region_label, role_label, status_label, provider_label, LANG_NAMES, ROLES, REGIONS,
)
from ..services import stats
from ..services.daily_report import build_report
from ..services.lang import lang_of, langs_of

logger = logging.getLogger(__name__)


class AdminStates(StatesGroup):
    search_user = State()
    add_participant = State()
    event_message = State()


# ─────────────────────────── доступ ───────────────────────────

async def is_admin(bot, tg_id: int) -> bool:
    if tg_id in bot["config"].tg_bot.admin_ids:
        return True
    return await sync_to_async(TGUser.objects.filter(tg_id=tg_id, is_admin=True).exists)()


# ─────────────────────────── клавиатуры ───────────────────────────

def main_kb() -> InlineKeyboardMarkup:
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(
        InlineKeyboardButton(t("adm_btn_stats"), callback_data="adm:stats"),
        InlineKeyboardButton(t("adm_btn_events"), callback_data="adm:events"),
    )
    kb.add(
        InlineKeyboardButton(t("adm_btn_find"), callback_data="adm:find"),
        InlineKeyboardButton(t("adm_btn_users_xlsx"), callback_data="adm:usersx"),
    )
    kb.add(InlineKeyboardButton(t("adm_btn_report"), callback_data="adm:rep"))
    kb.add(InlineKeyboardButton(t("adm_btn_bc"), callback_data="adm:bchelp"))
    return kb


def back_kb(to: str = "adm:menu", label_key: str = "adm_btn_menu") -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup().add(InlineKeyboardButton(t(label_key), callback_data=to))


# ─────────────────────────── поиск юзеров ───────────────────────────

from app_telegram.services import user_search_q as _user_query


@sync_to_async
def search_users(text: str, limit: int = 10):
    return list(TGUser.objects.filter(_user_query(text)).order_by('fullname')[:limit])


def _user_line(u: TGUser) -> str:
    parts = [u.fullname or "—"]
    if u.username:
        parts.append(f"@{u.username}")
    if u.region:
        parts.append(region_label(u.region))
    return " · ".join(parts)


# ─────────────────────────── вход ───────────────────────────

async def admin_entry(message: types.Message, state: FSMContext):
    if not await is_admin(message.bot, message.from_user.id):
        return
    await state.finish()
    await message.answer(t("adm_menu"), reply_markup=main_kb())


async def admin_callback(call: types.CallbackQuery, state: FSMContext):
    if not await is_admin(call.bot, call.from_user.id):
        await call.answer(t("no_access"), show_alert=True)
        return

    parts = call.data.split(":")
    action, args = parts[1], parts[2:]
    handler = CALLBACKS.get(action)
    if not handler:
        await call.answer()
        return
    try:
        await handler(call, state, *args)
    except exceptions.MessageNotModified:
        pass
    try:
        await call.answer()
    except exceptions.TelegramAPIError:
        pass  # уже ответили внутри хендлера


async def _edit_or_send(call: types.CallbackQuery, text: str, kb=None):
    try:
        await call.message.edit_text(text, reply_markup=kb, disable_web_page_preview=True)
    except (exceptions.MessageCantBeEdited, exceptions.BadRequest):
        await call.message.answer(text, reply_markup=kb, disable_web_page_preview=True)


async def cb_menu(call, state):
    await state.finish()
    await _edit_or_send(call, t("adm_menu"), main_kb())


async def cb_stats(call, state):
    await call.answer(t("adm_preparing"))
    await _edit_or_send(call, await build_report(), back_kb())


async def cb_bchelp(call, state):
    await _edit_or_send(call, t("adm_bc_help", regions=", ".join(REGIONS)), back_kb())


# ─────────────────────────── мероприятия ───────────────────────────

def _with_counts(qs):
    return qs.annotate(
        registered=Count('participants', filter=~Q(participants__status='rejected')),
        attended=Count('participants', filter=Q(participants__status='attended')),
    )


@sync_to_async
def _events(limit=15):
    return list(_with_counts(EcoProject.objects.all()).order_by('-date')[:limit])


async def cb_events(call, state):
    await state.finish()
    events = await _events()
    if not events:
        await _edit_or_send(call, t("adm_no_events"), back_kb())
        return
    kb = InlineKeyboardMarkup(row_width=1)
    for p in events:
        date = timezone.localtime(p.date).strftime('%d.%m')
        mark = "🟢" if p.is_active else "⚪️"
        kb.add(InlineKeyboardButton(
            f"{mark} {date} · {p.title[:30]} · {p.attended}/{p.registered}",
            callback_data=f"adm:ev:{p.id}",
        ))
    kb.add(InlineKeyboardButton(t("adm_btn_menu"), callback_data="adm:menu"))
    await _edit_or_send(call, t("adm_events_title"), kb)


@sync_to_async
def _event_card(project_id):
    return _with_counts(EcoProject.objects.filter(id=project_id)).first()


def event_kb(pid) -> InlineKeyboardMarkup:
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(
        InlineKeyboardButton(t("adm_btn_excel"), callback_data=f"adm:evx:{pid}"),
        InlineKeyboardButton(t("adm_btn_attended"), callback_data=f"adm:eva:{pid}"),
    )
    kb.add(InlineKeyboardButton(t("adm_btn_add"), callback_data=f"adm:evadd:{pid}"))
    kb.add(InlineKeyboardButton(t("adm_btn_msg"), callback_data=f"adm:evmsg:{pid}"))
    kb.add(InlineKeyboardButton(t("adm_btn_events_back"), callback_data="adm:events"))
    return kb


async def cb_event(call, state, pid):
    await state.finish()
    p = await _event_card(int(pid))
    if not p:
        await _edit_or_send(call, t("adm_event_not_found"), back_kb("adm:events", "adm_btn_events_back"))
        return
    text = t(
        "adm_event_card",
        title=escape(p.title),
        date=timezone.localtime(p.date).strftime('%d.%m.%Y %H:%M'),
        region=region_label(p.region),
        place=escape(p.location_name or '—'),
        active=t("adm_active") if p.is_active else t("adm_inactive"),
        reg=p.registered, max=p.max_participants, att=p.attended,
    )
    await _edit_or_send(call, text, event_kb(p.id))


@sync_to_async
def _participants(project_id, status=None):
    qs = ProjectParticipation.objects.filter(project_id=project_id).select_related('user', 'project')
    if status:
        qs = qs.filter(status=status)
    return list(qs.order_by('user__fullname'))


async def cb_event_attended(call, state, pid):
    parts = await _participants(int(pid), 'attended')
    if not parts:
        await call.message.answer(t("adm_no_attended"))
        return
    lines = [t("adm_attended_title", title=escape(parts[0].project.title), n=len(parts)), ""]
    lines += [f"{i}. {escape(_user_line(pp.user))}" for i, pp in enumerate(parts, 1)]
    text = "\n".join(lines)
    for x in range(0, len(text), 4000):
        await call.message.answer(text[x:x + 4000])


def _xlsx(title: str, headers: list, rows: list) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15803D")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _safe_filename(name: str) -> str:
    return re.sub(r"[^\w\-]+", "_", name, flags=re.UNICODE).strip("_")[:50] or "export"


async def cb_event_excel(call, state, pid):
    await call.answer(t("adm_preparing"))
    parts = [pp for pp in await _participants(int(pid)) if pp.status != 'rejected']
    if not parts:
        await call.message.answer(t("adm_no_participants"))
        return
    # сначала пришедшие — это и есть список на сертификаты
    parts.sort(key=lambda pp: (pp.status != 'attended', (pp.user.fullname or "").lower()))
    project = parts[0].project
    rows = []
    for i, pp in enumerate(parts, 1):
        u = pp.user
        rows.append([
            i, u.fullname, u.phone, u.email, f"@{u.username}" if u.username else "",
            region_label(u.region) if u.region else "", u.age, u.education_place,
            status_label(pp.status),
            timezone.localtime(pp.applied_at).strftime('%d.%m.%Y %H:%M') if pp.applied_at else "",
        ])
    buf = await sync_to_async(_xlsx)(t("xl_sheet_participants"), t("xl_event_headers"), rows)
    attended = sum(1 for pp in parts if pp.status == 'attended')
    date = timezone.localtime(project.date).strftime('%Y-%m-%d')
    await call.message.answer_document(
        types.InputFile(buf, filename=f"{date}_{_safe_filename(project.title)}.xlsx"),
        caption=t("adm_excel_caption", title=escape(project.title), att=attended, total=len(parts)),
    )


async def cb_users_excel(call, state):
    await call.answer(t("adm_preparing"))

    @sync_to_async
    def _rows():
        users = TGUser.objects.annotate(
            attended=Count('participations', filter=Q(participations__status='attended')),
        ).order_by('-created')
        return [
            [u.id, u.fullname, u.phone, u.email, f"@{u.username}" if u.username else "", u.tg_id,
             region_label(u.region) if u.region else "", u.age, role_label(u.role), u.balance,
             u.attended, provider_label(u.auth_provider), timezone.localtime(u.created).strftime('%d.%m.%Y')]
            for u in users.iterator(chunk_size=500)
        ]

    rows = await _rows()
    buf = await sync_to_async(_xlsx)(t("xl_sheet_users"), t("xl_user_headers"), rows)
    await call.message.answer_document(
        types.InputFile(buf, filename=f"users_{timezone.localdate().isoformat()}.xlsx"),
        caption=t("adm_users_caption", n=len(rows)),
    )


# ─────────────────────────── отчёт «кто пришёл» ───────────────────────────
# /admin → 📋 → период → регион → Excel (+ список в чат). Логика — app_telegram/reports.py,
# та же, что в Django-админке.

from app_telegram import reports


async def cb_report(call, state):
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(*[InlineKeyboardButton(t(f"rep_p_{p}"), callback_data=f"adm:repp:{p}") for p in reports.PERIODS])
    kb.add(InlineKeyboardButton(t("adm_btn_menu"), callback_data="adm:menu"))
    await _edit_or_send(call, t("rep_pick_period"), kb)


async def cb_report_period(call, state, period):
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(InlineKeyboardButton(t("rep_all_regions"), callback_data=f"adm:repr:{period}:all"))
    kb.add(InlineKeyboardButton("🏙 " + t("rep_tashkent"), callback_data=f"adm:repr:{period}:tashkent"))
    kb.add(*[InlineKeyboardButton(region_label(r), callback_data=f"adm:repr:{period}:{r}") for r in reports.REGION_CHOICES[2:]])
    kb.add(InlineKeyboardButton(t("btn_back"), callback_data="adm:rep"))
    await _edit_or_send(call, t("rep_pick_region", period=t(f"rep_p_{period}")), kb)


@sync_to_async
def _report_data(period, region):
    d_from, d_to = reports.period_range(period)
    regions = reports.region_codes(region)
    return d_from, d_to, reports.attendance(d_from, d_to, regions), reports.event_summary(d_from, d_to, regions)


async def cb_report_run(call, state, period, region):
    if period not in reports.PERIODS or region not in reports.REGION_CHOICES:
        return
    await call.answer(t("adm_preparing"))
    lang = await lang_of(call.from_user.id)
    d_from, d_to, parts, events = await _report_data(period, region)
    ptxt = f"{t(f'rep_p_{period}')} ({reports.period_text(d_from, d_to)})" if d_from else t("rep_p_all")
    rtxt = reports.region_choice_label(region)
    again = InlineKeyboardMarkup(row_width=1).add(
        InlineKeyboardButton(t("rep_btn_again"), callback_data="adm:rep"),
        InlineKeyboardButton(t("adm_btn_menu"), callback_data="adm:menu"),
    )
    if not parts:
        await call.message.answer(t("rep_empty", period=ptxt, region=escape(rtxt)), reply_markup=again)
        return
    title = f"{t('rep_sheet_people')} · {reports.period_text(d_from, d_to)} · {rtxt}"
    buf = await sync_to_async(reports.build_xlsx)(parts, events, lang, title)
    with_att = [e for e in events if e.attended]
    lines = "\n".join(f"• {escape(e.title)} — <b>{e.attended}</b>/{e.registered}" for e in with_att[:12])
    if len(with_att) > 12:
        lines += f"\n… +{len(with_att) - 12}"
    kb = InlineKeyboardMarkup(row_width=1).add(
        InlineKeyboardButton(t("rep_btn_text"), callback_data=f"adm:rept:{period}:{region}"),
        InlineKeyboardButton(t("rep_btn_again"), callback_data="adm:rep"),
    )
    caption = t("rep_caption", n=len(parts), people=len({pp.user_id for pp in parts}), events=len(with_att),
                period=ptxt, region=escape(rtxt), lines=lines)
    await call.message.answer_document(
        types.InputFile(buf, filename=reports.filename(d_from, d_to, region)),
        caption=caption[:1020], reply_markup=kb,
    )


async def cb_report_text(call, state, period, region):
    if period not in reports.PERIODS or region not in reports.REGION_CHOICES:
        return
    d_from, d_to, parts, _ = await _report_data(period, region)
    head = t("rep_list_title", period=t(f"rep_p_{period}"), region=escape(reports.region_choice_label(region)), n=len(parts))
    lines, cur = [head], None
    for i, pp in enumerate(parts, 1):
        if pp.project_id != cur:
            cur = pp.project_id
            lines += ["", f"📅 <b>{escape(pp.project.title)}</b>"]
        u = pp.user
        extra = " · ".join(x for x in (u.phone, f"@{u.username}" if u.username else "") if x)
        lines.append(f"{i}. {escape(u.fullname or '—')}" + (f" — {escape(extra)}" if extra else ""))
    text = "\n".join(lines)
    for x in range(0, len(text), 4000):
        await call.message.answer(text[x:x + 4000])


# ─────────────────────────── добавление участника ───────────────────────────

@sync_to_async
def add_to_event(user_id: int, project_id: int):
    """
    Добавляет юзера в мероприятие. Если мероприятие сегодня или уже прошло —
    сразу "attended" (+10 баллов через ProjectParticipation.save), иначе "approved".
    Возвращает (text, user, project, newly_attended).
    """
    user = TGUser.objects.get(id=user_id)
    project = EcoProject.objects.get(id=project_id)
    started = project.date <= timezone.now() or timezone.localtime(project.date).date() == timezone.localdate()
    status = 'attended' if started else 'approved'
    name = escape(user.fullname or "—")

    pp, created = ProjectParticipation.objects.get_or_create(user=user, project=project, defaults={'status': status})
    if created:
        key = "adm_added_attended" if status == 'attended' else "adm_added_registered"
        return t(key, name=name), user, project, status == 'attended'

    if pp.status != 'attended' and status == 'attended':
        pp.status = 'attended'
        pp.save()
        return t("adm_marked_attended", name=name), user, project, True

    return t("adm_already_in", name=name, status=status_label(pp.status)), user, project, False


async def _notify_attended(bot, user, project):
    if not user.tg_id:
        return
    user = await sync_to_async(TGUser.objects.get)(id=user.id)
    try:
        await bot.send_message(
            user.tg_id,
            t("attended_notify", await lang_of(user.tg_id), project=escape(project.title), balance=user.balance),
        )
    except Exception:
        pass


async def cb_event_add(call, state, pid):
    await AdminStates.add_participant.set()
    await state.update_data(project_id=int(pid))
    await call.message.answer(t("adm_add_prompt"), reply_markup=back_kb(f"adm:ev:{pid}", "adm_btn_done"))


async def add_participant_query(message: types.Message, state: FSMContext):
    if not await is_admin(message.bot, message.from_user.id):
        return
    pid = (await state.get_data()).get("project_id")
    users = await search_users(message.text or "")
    done_kb = back_kb(f"adm:ev:{pid}", "adm_btn_done")

    if not users:
        await message.answer(t("adm_not_found"), reply_markup=done_kb)
        return

    if len(users) == 1:
        text, user, project, newly = await add_to_event(users[0].id, pid)
        if newly:
            await _notify_attended(message.bot, user, project)
        await message.answer(text + t("adm_add_more"), reply_markup=done_kb)
        return

    kb = InlineKeyboardMarkup(row_width=1)
    for u in users:
        kb.add(InlineKeyboardButton(_user_line(u)[:60], callback_data=f"adm:addu:{pid}:{u.id}"))
    kb.add(InlineKeyboardButton(t("adm_btn_done"), callback_data=f"adm:ev:{pid}"))
    await message.answer(t("adm_many_found", n=len(users)), reply_markup=kb)


async def cb_add_user(call, state, pid, uid):
    text, user, project, newly = await add_to_event(int(uid), int(pid))
    if newly:
        await _notify_attended(call.bot, user, project)
    await call.message.answer(text)


# ─────────────────────────── сообщение участникам ───────────────────────────

async def cb_event_msg(call, state, pid):
    await AdminStates.event_message.set()
    await state.update_data(project_id=int(pid))
    await call.message.answer(t("adm_msg_prompt"), reply_markup=back_kb(f"adm:ev:{pid}", "adm_btn_cancel"))


async def event_message_send(message: types.Message, state: FSMContext):
    if not await is_admin(message.bot, message.from_user.id):
        return
    pid = (await state.get_data()).get("project_id")
    await state.finish()

    tg_ids = await sync_to_async(list)(
        ProjectParticipation.objects.filter(project_id=pid, user__tg_id__isnull=False)
        .exclude(status='rejected').values_list('user__tg_id', flat=True)
    )
    await message.answer(t("adm_sending", n=len(tg_ids)))
    sent = blocked = 0
    for tg_id in tg_ids:
        try:
            await message.copy_to(tg_id)
            sent += 1
        except exceptions.BotBlocked:
            blocked += 1
            await stats.mark_blocked(tg_id)
        except exceptions.RetryAfter as e:
            await asyncio.sleep(e.timeout)
        except Exception as e:
            logger.warning("event msg to %s failed: %s", tg_id, e)
        await asyncio.sleep(0.05)
    await message.answer(t("adm_sent", sent=sent, blocked=blocked),
                         reply_markup=back_kb(f"adm:ev:{pid}", "adm_btn_back_event"))


# ─────────────────────────── люди ───────────────────────────

async def cb_find(call, state):
    await AdminStates.search_user.set()
    await call.message.answer(t("adm_find_prompt"), reply_markup=back_kb("adm:menu", "adm_btn_cancel"))


async def search_user_query(message: types.Message, state: FSMContext):
    if not await is_admin(message.bot, message.from_user.id):
        return
    users = await search_users(message.text or "")
    if not users:
        await message.answer(t("adm_find_not_found"))
        return
    await state.finish()
    if len(users) == 1:
        text, kb = await _user_card(users[0].id)
        await message.answer(text, reply_markup=kb)
        return
    kb = InlineKeyboardMarkup(row_width=1)
    for u in users:
        kb.add(InlineKeyboardButton(_user_line(u)[:60], callback_data=f"adm:u:{u.id}"))
    kb.add(InlineKeyboardButton(t("adm_btn_menu"), callback_data="adm:menu"))
    await message.answer(t("adm_found_n", n=len(users)), reply_markup=kb)


@sync_to_async
def _user_data(user_id):
    u = TGUser.objects.get(id=user_id)
    parts = list(
        ProjectParticipation.objects.filter(user=u).select_related('project').order_by('-project__date')[:5]
    )
    counts = ProjectParticipation.objects.filter(user=u).aggregate(
        total=Count('id', filter=~Q(status='rejected')),
        att=Count('id', filter=Q(status='attended')),
    )
    return u, parts, counts['total'], counts['att']


async def _user_card(user_id):
    u, parts, total, attended = await _user_data(user_id)
    user_lang = await lang_of(u.tg_id) if u.tg_id else None
    text = t(
        "adm_user_card",
        name=escape(u.fullname or '—'),
        crown=" 👑" if u.is_admin else "",
        role=role_label(u.role),
        phone=escape(u.phone or '—'),
        email=escape(u.email or '—'),
        uname=f"@{escape(u.username)}" if u.username else "—",
        tg=u.tg_id or '—',
        region=region_label(u.region),
        age=u.age or '—',
        balance=u.balance, total=total, att=attended,
        provider=provider_label(u.auth_provider),
        date=timezone.localtime(u.created).strftime('%d.%m.%Y'),
        ulang=LANG_NAMES.get(user_lang, "—"),
    )
    if parts:
        text += "\n\n" + t("adm_last_events") + "\n"
        text += "\n".join(f"• {escape(pp.project.title)} — {status_label(pp.status)}" for pp in parts)

    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(
        InlineKeyboardButton(t("adm_btn_add_to_event"), callback_data=f"adm:uev:{u.id}"),
        InlineKeyboardButton(t("adm_btn_role"), callback_data=f"adm:ur:{u.id}"),
    )
    kb.add(InlineKeyboardButton(
        t("adm_btn_remove_admin") if u.is_admin else t("adm_btn_make_admin"), callback_data=f"adm:ua:{u.id}",
    ))
    kb.add(InlineKeyboardButton(t("adm_btn_menu"), callback_data="adm:menu"))
    return text, kb


async def cb_user(call, state, uid):
    await state.finish()
    text, kb = await _user_card(int(uid))
    await _edit_or_send(call, text, kb)


async def cb_user_roles(call, state, uid):
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(*[InlineKeyboardButton(role_label(code), callback_data=f"adm:urs:{uid}:{code}") for code in ROLES])
    kb.add(InlineKeyboardButton(t("btn_back"), callback_data=f"adm:u:{uid}"))
    await _edit_or_send(call, t("adm_choose_role"), kb)


async def cb_user_role_set(call, state, uid, role):
    if role not in TGUser.Role.values:
        return
    user = await sync_to_async(TGUser.objects.get)(id=int(uid))
    old_role = user.role
    user.role = role
    await sync_to_async(user.save)(update_fields=['role'])
    await call.answer(t("adm_saved"))

    # Как и в Django-админке: поздравляем только с повышением, не с демоцией в волонтёры.
    if user.tg_id and old_role != role and role != TGUser.Role.VOLUNTEER:
        lang = await lang_of(user.tg_id)
        try:
            await call.bot.send_message(user.tg_id, t("role_promo", lang, role=role_label(role, lang)))
        except Exception as e:
            logger.warning("role promo to %s failed: %s", user.tg_id, e)
    await cb_user(call, state, uid)


async def cb_user_admin(call, state, uid):
    user = await sync_to_async(TGUser.objects.get)(id=int(uid))
    if user.tg_id == call.from_user.id and user.is_admin:
        await call.answer(t("adm_cant_self"), show_alert=True)
        return
    user.is_admin = not user.is_admin
    await sync_to_async(user.save)(update_fields=['is_admin'])
    await call.answer(t("adm_saved"))
    await cb_user(call, state, uid)


async def cb_user_events(call, state, uid):
    events = await _events(10)
    kb = InlineKeyboardMarkup(row_width=1)
    for p in events:
        date = timezone.localtime(p.date).strftime('%d.%m')
        kb.add(InlineKeyboardButton(f"{date} · {p.title[:40]}", callback_data=f"adm:addu:{p.id}:{uid}"))
    kb.add(InlineKeyboardButton(t("btn_back"), callback_data=f"adm:u:{uid}"))
    await _edit_or_send(call, t("adm_choose_event"), kb)


# ─────────────────────────── блокировка бота ───────────────────────────

async def on_my_chat_member(update: types.ChatMemberUpdated):
    if update.chat.type != types.ChatType.PRIVATE:
        return
    status = update.new_chat_member.status
    if status == types.ChatMemberStatus.KICKED:
        await stats.mark_blocked(update.from_user.id)
    elif status == types.ChatMemberStatus.MEMBER:
        await stats.mark_unblocked(update.from_user.id)


CALLBACKS = {
    "menu": cb_menu,
    "stats": cb_stats,
    "bchelp": cb_bchelp,
    "events": cb_events,
    "ev": cb_event,
    "eva": cb_event_attended,
    "evx": cb_event_excel,
    "evadd": cb_event_add,
    "evmsg": cb_event_msg,
    "addu": cb_add_user,
    "usersx": cb_users_excel,
    "rep": cb_report,
    "repp": cb_report_period,
    "repr": cb_report_run,
    "rept": cb_report_text,
    "find": cb_find,
    "u": cb_user,
    "ur": cb_user_roles,
    "urs": cb_user_role_set,
    "ua": cb_user_admin,
    "uev": cb_user_events,
}


def register_admin_panel(dp: Dispatcher):
    dp.register_message_handler(admin_entry, commands=["admin"], state="*")
    dp.register_message_handler(admin_entry, text=variants("btn_admin"), state="*")
    dp.register_callback_query_handler(admin_callback, lambda c: c.data.startswith("adm:"), state="*")
    dp.register_message_handler(add_participant_query, state=AdminStates.add_participant)
    dp.register_message_handler(search_user_query, state=AdminStates.search_user)
    dp.register_message_handler(event_message_send, state=AdminStates.event_message, content_types=types.ContentTypes.ANY)
    dp.register_my_chat_member_handler(on_my_chat_member)
