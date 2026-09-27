"""
Отчёт «кто пришёл» — один на бот (/admin → 📋 Отчёт) и Django-админку
(Участия → 📥 Отчёт). Excel: лист 1 — люди, лист 2 — сводка по мероприятиям.

«Пришёл сегодня» = статус «пришёл» на мероприятии, которое было сегодня:
отдельного времени отметки в базе нет (applied_at — это время записи).
"""
from datetime import date, timedelta
from io import BytesIO

from django.db.models import Count, Q
from django.utils import timezone

from tgbot.i18n import REGIONS, region_label, t
from .models import EcoProject, ProjectParticipation
from .services import TASHKENT

PERIODS = ("today", "yesterday", "week", "month", "all")
# «tashkent» — город и область вместе (как везде в проекте)
REGION_CHOICES = ("all", "tashkent") + tuple(r for r in REGIONS if r not in TASHKENT)


def period_range(period: str):
    """(date_from, date_to) включительно; None — без ограничения."""
    today = timezone.localdate()
    if period == "today":
        return today, today
    if period == "yesterday":
        y = today - timedelta(days=1)
        return y, y
    if period == "week":
        return today - timedelta(days=6), today
    if period == "month":
        return today.replace(day=1), today
    return None, None


def region_codes(region: str):
    if region in (None, "", "all"):
        return None
    if region == "tashkent":
        return list(TASHKENT)
    return [region]


def region_choice_label(region: str, lang=None) -> str:
    if region == "all":
        return t("rep_all_regions", lang)
    if region == "tashkent":
        return t("rep_tashkent", lang)
    return region_label(region, lang)


def attendance(date_from: date | None, date_to: date | None, regions=None):
    """Отметки «пришёл» на мероприятиях за период в регионах (регион мероприятия)."""
    qs = ProjectParticipation.objects.filter(status='attended').select_related('user', 'project')
    if date_from:
        qs = qs.filter(project__date__date__gte=date_from)
    if date_to:
        qs = qs.filter(project__date__date__lte=date_to)
    if regions:
        qs = qs.filter(project__region__in=regions)
    return list(qs.order_by('project__date', 'project__title', 'user__fullname'))


def no_shows(date_from, date_to, regions=None):
    """Записались, но не пришли — только на уже прошедших мероприятиях."""
    qs = ProjectParticipation.objects.filter(status='approved', project__date__lte=timezone.now()).select_related('user', 'project')
    if date_from:
        qs = qs.filter(project__date__date__gte=date_from)
    if date_to:
        qs = qs.filter(project__date__date__lte=date_to)
    if regions:
        qs = qs.filter(project__region__in=regions)
    return list(qs.order_by('project__date', 'project__title', 'user__fullname'))


def event_summary(date_from, date_to, regions=None):
    qs = EcoProject.objects.all()
    if date_from:
        qs = qs.filter(date__date__gte=date_from)
    if date_to:
        qs = qs.filter(date__date__lte=date_to)
    if regions:
        qs = qs.filter(region__in=regions)
    return list(qs.annotate(
        registered=Count('participants', filter=~Q(participants__status='rejected')),
        attended=Count('participants', filter=Q(participants__status='attended')),
    ).order_by('date'))


def period_text(date_from, date_to, lang=None) -> str:
    if not date_from and not date_to:
        return t("rep_p_all", lang)
    f = date_from.strftime('%d.%m.%Y') if date_from else "…"
    to = date_to.strftime('%d.%m.%Y') if date_to else "…"
    return f if f == to else f"{f} — {to}"


def build_xlsx(parts, events, lang=None, title="", missed=None) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="15803D")

    def sheet(ws, headers, rows, caption):
        ws.append([caption]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(headers)
        for c in ws[3]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for col in ws.iter_cols(min_row=3):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 6), 45)
        ws.freeze_panes = "A4"
        if rows:
            ws.auto_filter.ref = f"A3:{ws.cell(row=3, column=len(headers)).column_letter}{3 + len(rows)}"

    wb = Workbook()
    ws = wb.active
    ws.title = t("rep_sheet_people", lang)[:31]
    rows = []
    for i, pp in enumerate(parts, 1):
        u, p = pp.user, pp.project
        rows.append([
            i, u.fullname, u.phone or "", f"@{u.username}" if u.username else "", u.email or "",
            u.age or "", u.education_place or "", region_label(u.region, lang) if u.region else "",
            p.title, region_label(p.region, lang) if p.region else "",
            timezone.localtime(p.date).strftime('%d.%m.%Y') if p.date else "", u.balance,
        ])
    sheet(ws, t("rep_people_headers", lang), rows, title)

    ws2 = wb.create_sheet(t("rep_sheet_events", lang)[:31])
    erows = []
    for i, e in enumerate(events, 1):
        pct = round(100 * e.attended / e.registered) if e.registered else 0
        erows.append([i, e.title, region_label(e.region, lang) if e.region else "",
                      timezone.localtime(e.date).strftime('%d.%m.%Y %H:%M') if e.date else "", e.registered, e.attended, f"{pct}%"])
    erows.append(["", t("rep_total", lang), "", "", sum(e.registered for e in events), sum(e.attended for e in events), ""])
    sheet(ws2, t("rep_event_headers", lang), erows, title)
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)

    if missed is not None:
        ws3 = wb.create_sheet(t("rep_sheet_noshow", lang)[:31])
        mrows = []
        for i, pp in enumerate(missed, 1):
            u, p = pp.user, pp.project
            mrows.append([i, u.fullname, u.phone or "", f"@{u.username}" if u.username else "", u.email or "",
                          u.age or "", u.education_place or "", region_label(u.region, lang) if u.region else "",
                          p.title, region_label(p.region, lang) if p.region else "",
                          timezone.localtime(p.date).strftime('%d.%m.%Y') if p.date else "", u.balance])
        sheet(ws3, t("rep_people_headers", lang), mrows, title)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def filename(date_from, date_to, region) -> str:
    span = "all" if not date_from else (date_from.isoformat() if date_from == date_to else f"{date_from.isoformat()}_{date_to.isoformat()}")
    return f"kelganlar_{region}_{span}.xlsx"
